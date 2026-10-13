"""Exportación a Excel: hojas de datos + una hoja de análisis por sección."""
from io import BytesIO
from typing import Dict, Optional
 
import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
 
from estilos import _grupos
from excel_graficas import mejorar_graficas
 
COLOR_A, COLOR_B, COLOR_C, COLOR_SIN = "3B6CF6", "12B886", "FF922B", "D5DBEA"
MARINO = "0B1646"
_borde = Side(style="thin", color="DCE5FA")
 
 
def _dar_formato_datos(ws) -> None:
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=MARINO)
        celda.alignment = Alignment(horizontal="center", vertical="center")
    for col in ws.columns:
        largo = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(largo + 3, 10), 40)
    ws.freeze_panes = "A2"
 
 
def _pintar(serie, color: str) -> None:
    serie.graphicalProperties.solidFill = color
    serie.graphicalProperties.line.solidFill = color
 
 
def _hoja_analisis(wb, nombre: str, sub: pd.DataFrame) -> None:
    ws = wb.create_sheet(f"Análisis {nombre}"[:31])
    ws["A1"] = f"Análisis por pregunta — {nombre}"
    ws["A1"].font = Font(bold=True, size=14, color=MARINO)
 
    for j, texto in enumerate(["Pregunta", "a", "b", "c", "Sin marca", "Total"], start=1):
        c = ws.cell(row=3, column=j, value=texto)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=MARINO)
        c.alignment = Alignment(horizontal="center")
 
    n = len(sub)
    for i, col in enumerate(sub.columns, start=4):
        s = sub[col]
        a, b, c_ = (int((s == o).sum()) for o in "abc")
        valores = [col.replace("Pregunta ", "P"), a, b, c_, n - a - b - c_, n]
        for j, v in enumerate(valores, start=1):
            celda = ws.cell(row=i, column=j, value=v)
            celda.border = Border(bottom=_borde)
            celda.alignment = Alignment(horizontal="left" if j == 1 else "center")
 
    fin = 3 + len(sub.columns)
    ws.column_dimensions["A"].width = 16
    for letra in "BCDEF":
        ws.column_dimensions[letra].width = 12
    cats = Reference(ws, min_col=1, min_row=4, max_row=fin)
 
    g1 = BarChart()
    g1.type, g1.grouping, g1.overlap = "col", "stacked", 100
    g1.title = "Respuestas por pregunta"
    g1.add_data(Reference(ws, min_col=2, max_col=4, min_row=3, max_row=fin), titles_from_data=True)
    g1.set_categories(cats)
    for serie, color in zip(g1.series, (COLOR_A, COLOR_B, COLOR_C)):
        _pintar(serie, color)
    ws.add_chart(g1, "A12")
 
    g2 = BarChart()
    g2.type, g2.grouping, g2.overlap = "bar", "percentStacked", 100
    g2.title = "Distribución por pregunta (100 %)"
    g2.add_data(Reference(ws, min_col=2, max_col=5, min_row=3, max_row=fin), titles_from_data=True)
    g2.set_categories(cats)
    for serie, color in zip(g2.series, (COLOR_A, COLOR_B, COLOR_C, COLOR_SIN)):
        _pintar(serie, color)
    g2.x_axis.scaling.orientation = "maxMin"
    g2.y_axis.number_format = "0%"
    ws.add_chart(g2, "A40")
 
 
def excel_bytes(hojas: Dict[str, pd.DataFrame], df: Optional[pd.DataFrame] = None) -> bytes:
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for nombre, tabla in hojas.items():
            tabla.to_excel(w, index=False, sheet_name=nombre[:31])
            _dar_formato_datos(w.sheets[nombre[:31]])
        if df is not None and len(df):
            for nombre, sub in _grupos(df).items():
                _hoja_analisis(w.book, nombre, sub)
            mejorar_graficas(w.book)
    return buf.getvalue()