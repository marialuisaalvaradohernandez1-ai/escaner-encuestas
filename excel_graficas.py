"""Mejora el diseño de las gráficas de Excel ya creadas con openpyxl.
 
Arregla: leyenda encima de las barras, título encimado con el eje, gráficas
pegadas entre sí, líneas de cuadrícula muy marcadas y celdas con cuadrícula.
 
Uso: llamar justo antes de guardar el libro.
    from excel_graficas import mejorar_graficas
    mejorar_graficas(w.book)      # w = pd.ExcelWriter(..., engine="openpyxl")
"""
import math
 
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
 
CM_POR_FILA = 0.53  # alto aproximado de una fila de Excel
 
 
def _linea_suave(eje) -> None:
    try:
        if eje.majorGridlines is not None:
            eje.majorGridlines.spPr = GraphicalProperties(
                ln=LineProperties(solidFill="E3E8F5", w=6350))
    except Exception:
        pass
 
 
def mejorar_graficas(wb, ancho_cm: float = 26.0) -> None:
    for ws in wb.worksheets:
        graficas = list(getattr(ws, "_charts", []))
        if not graficas:
            continue
 
        ws.sheet_view.showGridLines = False
        ancho_a = ws.column_dimensions["A"].width or 0
        ws.column_dimensions["A"].width = max(ancho_a, 16)
 
        n_filas_tabla = max(1, ws.max_row - 3)
        fila = ws.max_row + 3  # debajo de la tabla, con aire
 
        for ch in graficas:
            horizontal = getattr(ch, "type", "") == "bar"
            alto_cm = max(11.0, 1.0 * n_filas_tabla + 5) if horizontal else 11.0
            ch.width, ch.height = ancho_cm, alto_cm
 
            ch.layout = None
            if ch.legend is not None:
                ch.legend.position = "b"      # leyenda abajo, fuera de las barras
                ch.legend.overlay = False
            if ch.title is not None and hasattr(ch.title, "overlay"):
                ch.title.overlay = False      # título fuera del área de la gráfica
 
            try:
                ch.x_axis.delete = False
                ch.y_axis.delete = False
                ch.gapWidth = 60
                # Barras horizontales con categorías invertidas: el eje de % subía
                # y se encimaba con el título; así se queda abajo.
                if horizontal and ch.x_axis.scaling.orientation == "maxMin":
                    ch.y_axis.crosses = "max"
            except Exception:
                pass
            _linea_suave(ch.y_axis)
 
            ch.anchor = f"A{fila}"
            fila += math.ceil(alto_cm / CM_POR_FILA) + 3  # separación entre gráficas
 